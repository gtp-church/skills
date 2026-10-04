"""Packaging and behavior checks without requiring a presentation renderer."""
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import yaml
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / 'skills' / 'sermon-deck'
sys.path.insert(0, str(SKILL / 'scripts'))
import check
from deckbuilder import Deck


class PackagingTests(unittest.TestCase):
    def test_skill_manifests_and_resources(self):
        manifests = list((ROOT / 'skills').glob('*/SKILL.md'))
        self.assertTrue(manifests)
        for path in manifests:
            with self.subTest(skill=path.parent.name):
                match = re.match(r'^---\n(.*?)\n---\n', path.read_text(), re.S)
                self.assertIsNotNone(match)
                metadata = yaml.safe_load(match.group(1))
                self.assertEqual(metadata['name'], path.parent.name)
                self.assertRegex(metadata['name'], r'^[a-z0-9]+(?:-[a-z0-9]+)*$')
                self.assertTrue(metadata['description'].strip())
                self.assertLessEqual(len(metadata['description']), 1024)
                for resource in re.findall(r'`((?:scripts|assets|references)/[^`]+)`', path.read_text()):
                    self.assertTrue((path.parent / resource).exists(), resource)

    def test_example_builds_complete_deck(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / 'example.pptx'
            result = subprocess.run(
                [sys.executable, str(SKILL / 'references/example_build.py'), '--output', str(output)],
                cwd=tmp, text=True, capture_output=True, check=True)
            prs, problems = check.audit(output)
            self.assertEqual(problems, [])
            self.assertEqual(len(prs.slides), 11)
            self.assertIn('не попало в заметки: ничего', result.stdout)
            self.assertIn('продублировано: ничего', result.stdout)
            self.assertFalse((Path(tmp) / '_deck_logo.png').exists())

    def test_two_decks_can_be_built_independently(self):
        with tempfile.TemporaryDirectory() as tmp:
            first, second = Deck(), Deck()
            first.statement('ПЕРВАЯ', notes='Первый конспект')
            second.statement('ВТОРАЯ', notes='Второй конспект')
            for i, deck in enumerate((first, second)):
                path = Path(tmp) / f'{i}.pptx'
                deck.save(path)
                prs = Presentation(path)
                self.assertEqual(len(prs.slides), 1)
                self.assertEqual(sum(sh.shape_type == 13 for sh in prs.slides[0].shapes), 1)

    def test_libreoffice_resolution(self):
        with patch.dict('os.environ', {'SOFFICE': '/custom/soffice'}):
            self.assertEqual(check.find_soffice(), '/custom/soffice')
        with patch.dict('os.environ', {}, clear=True), patch('check.shutil.which', return_value='/usr/bin/soffice'):
            self.assertEqual(check.find_soffice(), '/usr/bin/soffice')
        with patch.dict('os.environ', {}, clear=True), patch('check.shutil.which', return_value=None):
            with patch('check.os.path.isfile', return_value=True):
                self.assertEqual(check.find_soffice(), '/Applications/LibreOffice.app/Contents/MacOS/soffice')
            with patch('check.os.path.isfile', return_value=False):
                with self.assertRaises(FileNotFoundError):
                    check.find_soffice()

    def test_failed_render_and_invalid_deck_fail_check(self):
        with patch('sys.argv', ['check.py', 'unused.pptx']), patch('check.package_problems', return_value=[]):
            with patch('check.audit', return_value=(Presentation(), [])), patch('check.render', side_effect=FileNotFoundError('missing renderer')):
                self.assertEqual(check.main(), 1)
            with tempfile.TemporaryDirectory() as tmp:
                with patch('check.audit', return_value=(Presentation(), ['missing notes'])), patch('check.render', return_value='test.pdf'), patch('check.audit_rendered_text', return_value=[]):
                    with patch('sys.argv', ['check.py', 'unused.pptx', tmp]):
                        self.assertEqual(check.main(), 1)


if __name__ == '__main__':
    unittest.main()
