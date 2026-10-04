"""Regression for the template defect that caused PowerPoint Repair.

Run: uv run --with python-pptx python -m unittest discover -s tests -v
"""
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from unittest.mock import patch

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / 'scripts'))

from deckbuilder import Deck
from packagecheck import package_problems
from pptx import Presentation


def corrupt_date_namespaces(source, destination):
    # Fault injection only: ElementTree rewrites namespace declarations while
    # leaving the dcterms QName in xsi:type unchanged, reproducing the defect.
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(destination, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == 'docProps/core.xml':
                data = ET.tostring(ET.fromstring(data), encoding='utf-8', xml_declaration=True)
            dst.writestr(info, data)


class PackageRegression(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name)
        self.template = SKILL / 'assets/template.pptx'
        self.broken = self.path / 'broken.pptx'
        corrupt_date_namespaces(self.template, self.broken)

    def test_template_and_standard_pptx_pass(self):
        self.assertEqual(package_problems(self.template), [])
        standard = self.path / 'standard.pptx'
        Presentation().save(standard)
        self.assertEqual(package_problems(standard), [])

    def test_fault_is_detected_even_though_python_pptx_can_read_it(self):
        Presentation(self.broken)
        errors = package_problems(self.broken)
        self.assertEqual(len(errors), 2, errors)
        self.assertTrue(all('docProps/core.xml' in error for error in errors))
        self.assertTrue(all('dcterms:W3CDTF' in error for error in errors))
        self.assertTrue(any('created' in error for error in errors))
        self.assertTrue(any('modified' in error for error in errors))

    def test_builder_rejects_bad_template(self):
        with self.assertRaisesRegex(ValueError, 'необъявленный префикс dcterms'):
            Deck(self.broken)

    def test_checker_rejects_fault_before_rendering(self):
        outdir = self.path / 'render'
        result = subprocess.run(
            [sys.executable, str(SKILL / 'scripts/check.py'), str(self.broken), str(outdir)],
            text=True, capture_output=True, check=False)
        self.assertEqual(result.returncode, 1)
        self.assertIn('необъявленный префикс dcterms', result.stderr)
        self.assertFalse(outdir.exists())

    def test_save_rechecks_metadata(self):
        deck = Deck()
        deck.title([('ПРОВЕРКА', 72, 'yellow')], notes='Заметка')
        original_save = deck.prs.save

        def save_then_corrupt(destination):
            original_save(destination)
            corrupted = self.path / 'corrupted-output.pptx'
            corrupt_date_namespaces(destination, corrupted)
            corrupted.replace(destination)

        with patch.object(deck.prs, 'save', side_effect=save_then_corrupt), self.assertRaisesRegex(
                ValueError, 'необъявленный префикс dcterms'):
            deck.save(self.path / 'bad-output.pptx')

    def test_all_slide_types_keep_notes_fonts_and_logo(self):
        deck = Deck()
        notes = [f'Фрагмент {i}: полный текст заметок.\nВторая строка.' for i in range(7)]
        deck.title([('ПРОВЕРКА', 72, 'yellow')], notes=notes[0])
        deck.divider('I', 'ПЕРВЫЙ РАЗДЕЛ', notes=notes[1])
        deck.bullets('Раздел', ['Пункт первый', 'Пункт второй'], notes=notes[2])
        deck.statement('БОГ ВЕРЕН', accent='ВЕРЕН', notes=notes[3])
        deck.callout('Важно', 'ПОЛНЫЙ ТЕКСТ', notes=notes[4])
        deck.verse('Римлянам 8:1', 'Итак нет ныне никакого осуждения', notes=notes[5])
        deck.poem(['Строка гимна'], 'Автор', notes=notes[6])
        destination = self.path / 'all-types.pptx'
        deck.save(destination)
        self.assertEqual(package_problems(destination), [])
        readback = Presentation(destination)
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in readback.slides], notes)
        with zipfile.ZipFile(self.template) as src, zipfile.ZipFile(destination) as dst:
            fonts = [name for name in src.namelist() if name.startswith('ppt/fonts/')]
            self.assertTrue(fonts)
            for name in fonts:
                self.assertEqual(src.read(name), dst.read(name), name)
            logos = {src.read(name) for name in src.namelist() if name.startswith('ppt/media/')}
            self.assertTrue(logos)
            for slide in readback.slides:
                pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
                self.assertEqual(len(pictures), 1)
                self.assertIn(pictures[0].image.blob, logos)


if __name__ == '__main__':
    unittest.main()
