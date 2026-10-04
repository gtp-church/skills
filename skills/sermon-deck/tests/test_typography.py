"""Observable regressions for readable Russian sermon slides."""
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / 'scripts'))

import check
from deckbuilder import AKROBAT, Deck
from textmetrics import text_width_pt, wrap_lines
from typography import display_text, trailing_preposition


class TypographyTests(unittest.TestCase):
    def test_narrow_wrap_moves_preposition_with_its_word(self):
        text = 'ВЕРУЮЩИЙ СОЕДИНЕН СО ХРИСТОМ'
        width = text_width_pt('ВЕРУЮЩИЙ СОЕДИНЕН СО', AKROBAT, 80)
        lines = wrap_lines(text, AKROBAT, 80, width)
        self.assertGreater(len(lines), 1)
        self.assertTrue(any('СО\u00a0ХРИСТОМ' in line for line in lines))
        self.assertFalse(any(trailing_preposition(line) for line in lines))
        self.assertEqual(' '.join(lines).replace('\u00a0', ' '), text)

    def test_protected_space_uses_normal_space_width(self):
        self.assertEqual(text_width_pt('СО ХРИСТОМ', AKROBAT, 60),
                         text_width_pt('СО\u00a0ХРИСТОМ', AKROBAT, 60))
        self.assertEqual(display_text('С Богом и во Христе'), 'С\u00a0Богом и во\u00a0Христе')
        self.assertEqual(display_text('пред Богом'), 'пред\u00a0Богом')

    def test_screen_changes_preserve_full_source_notes(self):
        source = 'Верующий соединён со Христом. Всё, что было Его, становится и её.'
        deck = Deck()
        deck.statement('ВЕРУЮЩИЙ СОЕДИНЁН\nСО ХРИСТОМ',
                       accent='СО ХРИСТОМ', notes=source)
        deck.divider('VIII', 'ЖИЗНЬ В СОЮЗЕ СО ХРИСТОМ', notes=source)
        deck.verse('Иоанна 1:12', 'А тем, которые приняли Его, дал власть быть чадами Божиими', notes=source)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'readable.pptx'
            deck.save(path)
            prs, problems = check.audit(path)
            self.assertEqual(problems, [])
            self.assertEqual([s.notes_slide.notes_text_frame.text for s in prs.slides], [source] * 3)
            for slide in prs.slides:
                text = '\n'.join(sh.text for sh in slide.shapes if sh.has_text_frame)
                self.assertNotIn('ё', text.lower())
            self.assertIn('СО\u00a0ХРИСТОМ', prs.slides[0].shapes[1].text)
            number = next(sh for sh in prs.slides[1].shapes if sh.has_text_frame and sh.text == 'VIII')
            self.assertLessEqual(number.text_frame.paragraphs[0].runs[0].font.size.pt, 64)
            quote = prs.slides[2]
            self.assertEqual(len([sh for sh in quote.shapes if sh.name == 'Quote decoration']), 2)
            ref = next(sh for sh in quote.shapes if sh.name == 'Quote reference')
            self.assertGreater(ref.top / 914400, 5)
            self.assertEqual(str(ref.text_frame.paragraphs[0].runs[0].font.color.rgb), 'FFFF00')

    def test_oversized_quote_requires_meaningful_split(self):
        with self.assertRaisesRegex(ValueError, 'раздели'):
            Deck().verse('Ссылка', 'Очень длинный текст ' * 100, notes='Источник')

    def test_headings_stay_compact_without_changing_body_leading(self):
        deck = Deck()
        text = 'ВЕРУЮЩИЙ СОЕДИНЕН\nСО ХРИСТОМ'
        compact = deck.statement(text, notes='Источник')
        spaced = deck.statement(text, notes='Источник', leading=1.05)
        divider = deck.divider('I', 'ПРАВЕДНОСТЬ\nВ ВЕТХОМ ЗАВЕТЕ', notes='Источник')
        title = deck.title([('ПЕРВАЯ', 96, 'yellow'), ('ВТОРАЯ', 96, 'white')], notes='Источник')
        verse = deck.verse('Римлянам 8:1', 'Итак нет ныне никакого осуждения', notes='Источник')
        compact_box, spaced_box = compact.shapes[1], spaced.shapes[1]
        self.assertLess(compact_box.height, spaced_box.height)
        self.assertAlmostEqual(compact_box.top + compact_box.height / 2,
                               spaced_box.top + spaced_box.height / 2, delta=1)
        for box in (compact_box, divider.shapes[1], title.shapes[2]):
            self.assertTrue(all(p.line_spacing < 1 for p in box.text_frame.paragraphs))
        quote_text = next(sh for sh in verse.shapes if sh.name == 'Quote text')
        self.assertEqual(quote_text.text_frame.paragraphs[0].line_spacing, 1.05)

    def test_both_angular_quote_strokes_survive_save_without_a_font(self):
        deck = Deck()
        deck.verse('Иоанна 1:12', 'А тем, которые приняли Его', notes='Источник')
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'quotes.pptx'
            deck.save(path)
            slide = Presentation(path).slides[0]
            decorations = [sh for sh in slide.shapes if sh.name == 'Quote decoration']
            self.assertEqual(len(decorations), 2)
            for sh in decorations:
                self.assertFalse(sh.text.strip())
                self.assertEqual(len(sh._element.xpath('.//a:custGeom//a:moveTo')), 2)
                self.assertEqual(len(sh._element.xpath('.//a:custGeom//a:close')), 2)
                self.assertEqual(len(sh._element.xpath('.//a:custGeom//a:lnTo')), 6)
                self.assertFalse(sh._element.xpath('.//a:cubicBezTo | .//a:quadBezTo | .//a:arcTo'))
                self.assertGreaterEqual(sh.left, 0)
                self.assertGreaterEqual(sh.top, 0)
                self.assertLessEqual(sh.left + sh.width, deck.prs.slide_width)
                self.assertLessEqual(sh.top + sh.height, deck.prs.slide_height)
                self.assertEqual(str(sh.fill.fore_color.rgb), '3D3D3D')
            text_box = next(sh for sh in slide.shapes if sh.name == 'Quote text')
            self.assertTrue(all(slide.shapes.index(sh) < slide.shapes.index(text_box)
                                for sh in decorations))

    def test_real_rendered_line_endings_are_checked(self):
        xml = '''<html xmlns="http://www.w3.org/1999/xhtml"><body><doc>
          <page><flow><block><line><word>ЖИЗНЬ</word><word>СО</word></line>
          <line><word>ХРИСТОМ</word></line></block></flow></page>
          <page><flow><block><line><word>ЖИЗНЬ</word></line>
          <line><word>СО</word><word>ХРИСТОМ</word></line></block></flow></page>
        </doc></body></html>'''
        with patch('check.subprocess.run') as run:
            run.return_value.stdout = xml
            problems = check.audit_rendered_text('rendered.pdf')
        self.assertEqual(len(problems), 1)
        self.assertIn('слайд 1', problems[0])


if __name__ == '__main__':
    unittest.main()
