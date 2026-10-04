# -*- coding: utf-8 -*-
"""Сборка презентации проповеди в фирменном стиле.

Всё делается через API python-pptx: библиотека сама собирает пакет
(Content_Types, связи, presentation.xml). Писать XML слайдов строками нельзя.
Шаблон и результат проверяются также на ошибки пакета: API сохраняет
повреждённые метаданные шаблона, включая необъявленные QName-префиксы.

Пример:

    from deckbuilder import Deck
    d = Deck()                     # вшитый шаблон; или Deck(newest_template(папка))
    d.title([('ИСКУПЛЕНИЕ', 150, 'yellow'), ('В ЧЕМ СМЫСЛ?', 72, 'white')],
            notes='...')
    d.divider('I', 'ПРИРОДА БОГА – СВЯТОСТЬ', notes='...')
    d.bullets('Природа Бога – святость', ['Святость – сущность Бога'], notes='...')
    d.save('Проповедь.pptx', title='В чем смысл Искупления?')
"""
import os
from io import BytesIO
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from textmetrics import longest_word_pt, wrap_lines
from packagecheck import require_valid_package
from typography import display_text

EMU = 914400
PT = EMU / 72
LINE = 1.213          # множитель высоты строки, выверен по слайдам шаблона
LEADING = 1.05        # запас между строками для нижних выносных элементов
INSETS = 7.2          # верхнее + нижнее поле надписи, пункты

AKROBAT, HEADING, HEADING_XB, INTRO = (
    'Akrobat Bold', 'Heading Pro', 'Heading Pro ExtBd', 'Intro ')
YELLOW = RGBColor(0xFF, 0xFF, 0x00)
WHITE = MSO_THEME_COLOR.BACKGROUND_1
COLORS = {'yellow': YELLOW, 'white': WHITE}

SLIDE_W, SLIDE_H = 12192000, 6858000
LOGO_POS, LOGO_SZ = (11508153, 6131169), 541871
HEAD_X, HEAD_CX = 1028125, 10480028          # рамка заголовков
BODY_X, BODY_CX = 1022693, 10508225          # рамка списков и текста
EYE_X, EYE_Y, EYE_CX = 645485, 474510, 11169298   # жёлтая плашка раздела
MARL = 800100                                 # отступ под номер пункта
CIRCLE = (1201598, 492852, 2239836, 2252997)  # круг с римской цифрой
NUM_X, NUM_CX = 1386285, 1870516


# --- расчёт размеров: подбираем самый крупный кегль, который влезает --------
def avail_pt(cx, marl=0):
    return (cx - 91440 * 2 - marl) / PT


def nlines(text, font, size, cx, marl=0):
    return len(wrap_lines(text, font, size, avail_pt(cx, marl)))


def word_fits(text, font, size, cx, marl=0):
    """Слово шире строки PowerPoint разрывает посередине — так нельзя."""
    return longest_word_pt(text, font, size) <= avail_pt(cx, marl)


def fit_size(text, font, cx, candidates, max_lines, marl=0):
    for s in candidates:
        if word_fits(text, font, s, cx, marl) and nlines(text, font, s, cx, marl) <= max_lines:
            return s
    return candidates[-1]


def block_h(paras):
    """paras: [(строк, кегль, межстрочный, отступ после)] -> высота в EMU."""
    total = 0.0
    for i, (n, size, lnsp, aft) in enumerate(paras):
        total += n * size * LINE * lnsp
        if i < len(paras) - 1:
            total += aft
    return int((total + INSETS) / 72 * EMU)


BUNDLED_TEMPLATE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), '..', 'assets', 'template.pptx')


def newest_template(folder=None):
    """Шаблон, задающий стиль.

    Без аргумента — вшитый в скилл `assets/template.pptx` (та самая «Пасха 2026»
    со всеми вшитыми шрифтами и логотипом). Скилл самодостаточен и не зависит
    от папки с готовыми презентациями.

    Если папка с референсами под рукой — передай её, и возьмётся самый свежий
    файл: стиль общины со временем меняется, и свежая презентация точнее.
    Даты записей в zip у этих файлов затёрты (1980), поэтому смотрим
    dcterms:modified в docProps/core.xml.
    """
    if folder is None:
        return os.path.normpath(BUNDLED_TEMPLATE)
    import re
    best, best_stamp = None, ''
    for name in sorted(os.listdir(folder)):
        if not name.endswith('.pptx') or name.startswith('~'):
            continue
        path = os.path.join(folder, name)
        try:
            with zipfile.ZipFile(path) as z:
                core = z.read('docProps/core.xml').decode('utf-8')
            stamp = re.search(r'<dcterms:modified[^>]*>([^<]+)<', core).group(1)
        except Exception:
            continue
        if stamp > best_stamp:
            best, best_stamp = path, stamp
    return best or os.path.normpath(BUNDLED_TEMPLATE)


class Deck:
    def __init__(self, template=None):
        template = template or newest_template()
        require_valid_package(template)
        self.template = template
        self.prs = Presentation(template)
        for sld in list(self.prs.slides._sldIdLst):        # убрать слайды шаблона
            self.prs.part.drop_rel(sld.get(qn('r:id')))
            self.prs.slides._sldIdLst.remove(sld)
        self.blank = self.prs.slide_masters[0].slide_layouts[6]
        with zipfile.ZipFile(template) as z:
            media = [n for n in z.namelist() if n.startswith('ppt/media/')]
            self._logo = z.read(sorted(media)[0])

    # --- примитивы ---------------------------------------------------------
    def _slide(self):
        s = self.prs.slides.add_slide(self.blank)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        sh.name = 'Rectangle 9'
        sh.fill.solid()
        sh.fill.fore_color.theme_color = MSO_THEME_COLOR.TEXT_1   # чёрный фон
        sh.line.fill.background()
        sh.shadow.inherit = False
        return s

    def _logo_on(self, slide):
        slide.shapes.add_picture(BytesIO(self._logo), LOGO_POS[0], LOGO_POS[1], LOGO_SZ, LOGO_SZ)

    def _box(self, slide, x, y, cx, cy, name='Rectangle 14'):
        sh = slide.shapes.add_textbox(int(x), int(y), int(cx), int(cy))
        sh.name = name
        tf = sh.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
        return tf

    @staticmethod
    def _run(p, text, font, size, color, bold=False):
        text = display_text(text)
        if '\n' in text:
            chunks = text.split('\n')
            for i, chunk in enumerate(chunks):
                if i:
                    p.add_line_break()
                r = Deck._run(p, chunk, font, size, color, bold)
            return r
        r = p.add_run()
        r.text = text
        r.font.name = font
        r.font.size = Pt(size)
        if bold:
            r.font.bold = True
        col = COLORS.get(color, color)
        if isinstance(col, RGBColor):
            r.font.color.rgb = col
        else:
            r.font.color.theme_color = col
        return r

    @staticmethod
    def _numbered(p):
        """Автонумерация «1.» с висячим отступом, как в шаблоне.
        Порядок детей a:pPr задан схемой: lnSpc, spcBef, spcAft, buClr, buSzPct, buAutoNum."""
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(MARL))
        pPr.set('indent', str(-MARL))
        clr = pPr.makeelement(qn('a:buClr'), {})
        clr.append(clr.makeelement(qn('a:schemeClr'), {'val': 'bg1'}))
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn('a:buSzPct'), {'val': '80000'}))
        pPr.append(pPr.makeelement(qn('a:buAutoNum'), {'type': 'arabicPeriod'}))

    def _para(self, tf, first, align=None, lnsp=None, aft=None, number=False):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        if lnsp is not None:
            p.line_spacing = lnsp
        if aft is not None:
            p.space_after = Pt(aft)
        if number:
            self._numbered(p)
        return p

    def _eyebrow(self, slide, text):
        text = text.upper()
        size = 48
        while size > 28 and nlines('>  ' + text, HEADING, size, EYE_CX) > 1:
            size -= 4
        tf = self._box(slide, EYE_X, EYE_Y, EYE_CX, block_h([(1, size, 1.0, 0)]), 'TextBox 5')
        p = tf.paragraphs[0]
        self._run(p, '>', HEADING_XB, size, YELLOW, bold=True)
        self._run(p, ' ', HEADING, size, YELLOW)
        self._run(p, text, HEADING, size, YELLOW)

    def _notes(self, slide, notes):
        tf = slide.notes_slide.notes_text_frame
        tf.clear()
        for j, blk in enumerate(b for b in (notes or '').split('\n\n') if b.strip()):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = blk.strip()
            r.font.size = Pt(14)

    # --- типы слайдов ------------------------------------------------------
    def title(self, lines, notes='', font=HEADING):
        """lines: [(текст, кегль, 'yellow'|'white')] — по строке на абзац."""
        s = self._slide()
        self._logo_on(s)
        cy = block_h([(1, sz, LEADING, 0) for _, sz, _ in lines])
        tf = self._box(s, 0, (SLIDE_H - cy) // 2, SLIDE_W - 1, cy, 'TextBox 3')
        for i, (t, sz, col) in enumerate(lines):
            p = self._para(tf, i == 0, align=PP_ALIGN.CENTER, lnsp=LEADING)
            self._run(p, t, font, sz, col)
        self._notes(s, notes)
        return s

    def divider(self, numeral, heading, notes='', center_in=4.80):
        """Разделитель: римская цифра в жёлтом круге + крупный заголовок."""
        s = self._slide()
        size = fit_size(heading, AKROBAT, HEAD_CX, [115, 105, 96, 88, 80], 2)
        cy = block_h([(nlines(heading, AKROBAT, size, HEAD_CX), size, LEADING, 0)])
        tf = self._box(s, HEAD_X, center_in * EMU - cy / 2, HEAD_CX, cy)
        self._run(self._para(tf, True, lnsp=LEADING), heading, AKROBAT, size, WHITE, bold=True)

        nsize = next((n for n in (115, 105, 96, 88, 80, 72, 64, 60)
                      if longest_word_pt(numeral, INTRO, n) <= avail_pt(NUM_CX)), 60)
        nh = block_h([(1, nsize, 1.0, 0)])
        ntf = self._box(s, NUM_X, CIRCLE[1] + CIRCLE[3] / 2 - nh / 2, NUM_CX, nh, 'TextBox 3')
        self._run(self._para(ntf, True, align=PP_ALIGN.CENTER), numeral, INTRO, nsize, YELLOW)

        oval = s.shapes.add_shape(MSO_SHAPE.OVAL, *CIRCLE)
        oval.name = 'Oval 4'
        oval.fill.background()
        oval.line.color.rgb = YELLOW
        oval.line.width = Emu(38100)
        oval.shadow.inherit = False
        self._logo_on(s)
        self._notes(s, notes)
        return s

    def bullets(self, section, items, notes='', size=None, top_in=None,
                center_in=4.35, all_items=None):
        """Нумерованный список под плашкой раздела.

        all_items — полный перечень: передавай его на «нарастающих» слайдах,
        чтобы кегль и положение блока не прыгали от слайда к слайду.
        Вместе с top_in блок прижимается к верху и растёт вниз.
        """
        s = self._slide()
        ref = all_items or items
        aft = 24 if len(ref) <= 4 else 12
        limit = 4.95 if top_in is not None else 4.85
        if size is None:
            size = 44
            for cand in (60, 54, 48, 44):
                nls = [nlines(t, AKROBAT, cand, BODY_CX, MARL) for t in ref]
                h = block_h([(n, cand, LEADING, aft) for n in nls]) / EMU
                if (max(nls) <= 2 and h <= limit
                        and all(word_fits(t, AKROBAT, cand, BODY_CX, MARL) for t in ref)):
                    size = cand
                    break
        cy = block_h([(nlines(t, AKROBAT, size, BODY_CX, MARL), size, LEADING, aft) for t in items])
        y = top_in * EMU if top_in is not None else center_in * EMU - cy / 2
        tf = self._box(s, BODY_X, y, BODY_CX, cy)
        for i, t in enumerate(items):
            p = self._para(tf, i == 0, lnsp=LEADING, aft=aft, number=True)
            self._run(p, t, AKROBAT, size, WHITE, bold=True)
        self._logo_on(s)
        self._eyebrow(s, section)
        self._notes(s, notes)
        return size

    def statement(self, text, accent=None, notes='', center_in=3.75):
        """Одна крупная мысль на весь слайд; accent — часть строки жёлтым."""
        s = self._slide()
        text = display_text(text)
        accent = display_text(accent) if accent else None
        size = fit_size(text, AKROBAT, HEAD_CX, [115, 105, 96, 88, 80, 72], 3)
        cy = block_h([(nlines(text, AKROBAT, size, HEAD_CX), size, LEADING, 0)])
        tf = self._box(s, HEAD_X, center_in * EMU - cy / 2, HEAD_CX, cy)
        p = self._para(tf, True, lnsp=LEADING)
        if accent and accent in text:
            i = text.index(accent)
            for chunk, col in ((text[:i], WHITE), (accent, YELLOW), (text[i + len(accent):], WHITE)):
                if chunk:
                    self._run(p, chunk, AKROBAT, size, col, bold=True)
        else:
            self._run(p, text, AKROBAT, size, WHITE, bold=True)
        self._logo_on(s)
        self._notes(s, notes)
        return s

    def callout(self, section, text, notes='', center_in=4.15,
                sizes=(72, 66, 60, 54, 48)):
        """Плашка («ВАЖНО!», «ОСТОРОЖНО!») и развёрнутая мысль под ней."""
        s = self._slide()
        size = fit_size(text, AKROBAT, BODY_CX, list(sizes), 4)
        cy = block_h([(nlines(text, AKROBAT, size, BODY_CX), size, LEADING, 0)])
        tf = self._box(s, BODY_X, center_in * EMU - cy / 2, BODY_CX, cy)
        self._run(self._para(tf, True, lnsp=LEADING), text, AKROBAT, size, WHITE, bold=True)
        self._logo_on(s)
        self._eyebrow(s, section)
        self._notes(s, notes)
        return s

    def verse(self, reference, text, notes=''):
        """Цитата: серые кавычки на фоне, белый текст, жёлтая ссылка внизу."""
        s = self._slide()
        # Editable glyphs from an already embedded font. Their intentional crop
        # is decorative; the quote text and its reference stay inside the slide.
        for x, y in ((-2.846, -2.356), (7.059, 0.976)):
            tf = self._box(s, x * EMU, y * EMU, 6 * EMU, 12 * EMU, 'Quote decoration')
            tf.auto_size = MSO_AUTO_SIZE.NONE
            self._run(tf.paragraphs[0], '”', 'Calibri', 857, RGBColor(0x3D, 0x3D, 0x3D))
        text = text.strip().removeprefix('«').removesuffix('»')
        size = next((candidate for candidate in (60, 54, 48, 44, 40)
                     if word_fits(text, AKROBAT, candidate, BODY_CX)
                     and nlines(text, AKROBAT, candidate, BODY_CX) <= 5
                     and block_h([(nlines(text, AKROBAT, candidate, BODY_CX),
                                   candidate, LEADING, 0)]) <= 4.6 * EMU), None)
        if size is None:
            raise ValueError('Цитата слишком длинная: раздели её по смыслу на два слайда')
        cy = block_h([(nlines(text, AKROBAT, size, BODY_CX), size, LEADING, 0)])
        tf = self._box(s, BODY_X, 3.3 * EMU - cy / 2, BODY_CX, cy, 'Quote text')
        self._run(self._para(tf, True, lnsp=LEADING), text, AKROBAT, size, WHITE, bold=True)
        tf = self._box(s, 3.0 * EMU, 5.55 * EMU, 8.9 * EMU,
                       block_h([(1, 44, LEADING, 0)]), 'Quote reference')
        self._run(self._para(tf, True, align=PP_ALIGN.RIGHT, lnsp=LEADING),
                  reference, AKROBAT, 44, YELLOW, bold=True)
        self._logo_on(s)
        self._notes(s, notes)
        return s

    def poem(self, lines, attribution, notes=''):
        """Куплет гимна: строки по центру, ниже — автор."""
        s = self._slide()
        px, pcx = 457200, 11277600
        size = 48
        for cand in (60, 54, 48, 44, 40):
            if all(nlines(t, AKROBAT, cand, pcx) == 1 and word_fits(t, AKROBAT, cand, pcx)
                   for t in lines) and block_h([(1, cand, 0.9, 0)] * len(lines)) / EMU <= 5.15:
                size = cand
                break
        asize, gap = 36, 28
        cy = block_h([(1, size, 0.9, 0)] * len(lines) + [(1, gap, 0.9, 0), (1, asize, 0.9, 0)])
        tf = self._box(s, px, 3.6 * EMU - cy / 2, pcx, cy)
        for i, t in enumerate(lines):
            self._run(self._para(tf, i == 0, align=PP_ALIGN.CENTER, lnsp=0.9),
                      t, AKROBAT, size, WHITE, bold=True)
        self._para(tf, False, align=PP_ALIGN.CENTER, lnsp=0.9).add_run().font.size = Pt(gap)
        self._run(self._para(tf, False, align=PP_ALIGN.CENTER, lnsp=0.9),
                  attribution, HEADING, asize, YELLOW)
        self._logo_on(s)
        self._notes(s, notes)
        return s

    def save(self, path, title=None):
        if title:
            self.prs.core_properties.title = title
        self.prs.core_properties.revision = 1
        if os.path.exists(path):
            os.remove(path)
        self.prs.save(path)
        require_valid_package(path)
        return path
